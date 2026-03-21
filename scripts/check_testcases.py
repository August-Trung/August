from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    path = Path("TEST_CASES.xlsx")
    if not path.exists():
        print("TEST_CASES.xlsx not found.")
        return 1

    wb = load_workbook(path)
    failures: list[tuple[str, str, str]] = []
    pending: list[tuple[str, str, str]] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            case_id = str(row[0] or "").strip()
            status = str(row[7] or "").strip().lower()
            notes = str(row[11] or "").strip()
            if status == "fail":
                failures.append((sheet.title, case_id, notes))
            elif status == "not run":
                pending.append((sheet.title, case_id, notes))

    if failures:
        print("Found failed test cases:")
        for sheet, case_id, notes in failures:
            print(f"- [{sheet}] {case_id}: {notes}")
    if pending:
        print("Found NOT RUN test cases (must be executed first):")
        for sheet, case_id, notes in pending:
            print(f"- [{sheet}] {case_id}: {notes}")
    if failures or pending:
        return 2

    print("No failed or not-run test case rows in TEST_CASES.xlsx.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
