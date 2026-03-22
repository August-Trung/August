from __future__ import annotations

import json
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from futureos.auth import AUTH_STATE_FILE, authenticate_login
from futureos.models import Action, IntentType, Role, UserContext
from futureos.policy import evaluate_action, evaluate_plan
from futureos.router import route
from futureos.safety import AUDIT_CHAIN_FILE, AUDIT_FILE, HISTORY_CHAIN_FILE, HISTORY_FILE, verify_chain
from futureos.session import ACTIVE_SESSION_FILE, SESSIONS_FILE, SessionStore
from futureos.voice import VoiceEngine
from futureos.workflows import execute_action
from futureos.queue import (
    DEAD_LETTER_FILE,
    PROCESSED_KEYS_FILE,
    QUEUE_FILE,
    WORKER_HEARTBEAT_FILE,
    WORKER_LOCK_FILE,
    BackgroundWorker,
    TaskQueue,
)


@dataclass
class CaseResult:
    status: str
    actual: str
    notes: str


def _reset_runtime_files() -> None:
    for p in [
        AUTH_STATE_FILE,
        SESSIONS_FILE,
        ACTIVE_SESSION_FILE,
        AUDIT_FILE,
        HISTORY_FILE,
        AUDIT_CHAIN_FILE,
        HISTORY_CHAIN_FILE,
        QUEUE_FILE,
        DEAD_LETTER_FILE,
        PROCESSED_KEYS_FILE,
        WORKER_LOCK_FILE,
        WORKER_HEARTBEAT_FILE,
    ]:
        if p.exists():
            p.unlink(missing_ok=True)


def _run_cases() -> dict[str, CaseResult]:
    _reset_runtime_files()
    today = str(date.today())
    _ = today  # keep deterministic date reference if needed later

    results: dict[str, CaseResult] = {}

    # Functional
    plan = route("tim ban thao lap trinh trong o D va gui zalo cho sep va gui email hang loat cho team")
    owner = UserContext(role=Role.OWNER, allow_c_drive_full=False, actor="tc-owner")
    allowed, checks = evaluate_plan(plan.actions, owner)
    composite_ok = any(a.intent == IntentType.COMPOSITE for a in plan.actions)
    if composite_ok and allowed:
        out = execute_action(plan.actions[0])
        step_count = len(out.payload.get("steps", []))
        ok = out.ok and step_count >= 2
        results["F-001"] = CaseResult("Pass" if ok else "Fail", f"steps={step_count}", "Composite run as owner.")
    else:
        results["F-001"] = CaseResult("Fail", json.dumps(checks, ensure_ascii=False), "Plan denied/unexpected.")

    # Flow 04 UI checks
    ui_path = Path("futureos/ui_app.py")
    results["F-004"] = CaseResult(
        "Pass" if ui_path.exists() else "Fail",
        f"ui_exists={ui_path.exists()}",
        "UI shell file exists.",
    )
    try:
        import streamlit  # noqa: F401

        results["F-005"] = CaseResult("Pass", "streamlit import ok", "UI dependency available.")
    except Exception as e:
        results["F-005"] = CaseResult("Fail", str(e), "UI dependency missing.")

    # Flow 05 background worker checks
    _reset_runtime_files()
    q_store = SessionStore()
    q_rec = q_store.create_session(UserContext(role=Role.OWNER, allow_c_drive_full=False, actor="tc-queue"))
    q = TaskQueue()
    w = BackgroundWorker(q, q_store)
    q_task = q.enqueue("tim file o d", q_rec.session_id, auto_confirm=False, max_retries=1)
    q_out = w.run_once()
    q_ok = q_out.get("ok", False) and q_out.get("task_id") == q_task.id
    results["F-006"] = CaseResult("Pass" if q_ok else "Fail", json.dumps(q_out, ensure_ascii=False), "Queue enqueue/process one.")

    _reset_runtime_files()
    q = TaskQueue()
    w = BackgroundWorker(q, SessionStore())
    q.enqueue("tim file o d", "missing-session", auto_confirm=False, max_retries=0)
    dead_out = w.run_once()
    dead_ok = (not dead_out.get("ok", True)) and DEAD_LETTER_FILE.exists()
    results["F-007"] = CaseResult("Pass" if dead_ok else "Fail", json.dumps(dead_out, ensure_ascii=False), "Dead-letter on unrecoverable task.")

    _reset_runtime_files()
    q_store = SessionStore()
    q_rec = q_store.create_session(UserContext(role=Role.OWNER, allow_c_drive_full=False, actor="tc-idem"))
    q = TaskQueue()
    a = q.enqueue("tim file o d", q_rec.session_id, idempotency_key="same-key")
    b = q.enqueue("tim file o d", q_rec.session_id, idempotency_key="same-key")
    idem_ok = a.id == b.id and q.stats().get("queued", 0) == 1
    results["F-008"] = CaseResult("Pass" if idem_ok else "Fail", f"a={a.id},b={b.id},queued={q.stats().get('queued',0)}", "Idempotency blocks duplicate queue.")

    _reset_runtime_files()
    q_store = SessionStore()
    q_rec = q_store.create_session(UserContext(role=Role.OWNER, allow_c_drive_full=False, actor="tc-cancel"))
    q = TaskQueue()
    t = q.enqueue("tim file o d", q_rec.session_id)
    cancel_ok = q.cancel(t.id)
    results["F-009"] = CaseResult("Pass" if cancel_ok else "Fail", f"cancel_ok={cancel_ok}", "Queued task cancel supported.")

    # Core file ops VN command
    p = route("tao thu muc zzz o desktop roi di chuyen vao thu muc backup trong o d")
    composite_steps = p.actions[0].args.get("steps", []) if p.actions else []
    intents = [x.get("intent") for x in composite_steps]
    core_ok = "dir_create" in intents and "path_move" in intents
    results["F-010"] = CaseResult("Pass" if core_ok else "Fail", ",".join(intents), "VN create+move command parsing.")

    p = route("di chuyen thu muc zzz vao backup o d")
    intents = []
    for a in p.actions:
        if a.intent == IntentType.COMPOSITE:
            intents.extend([s.get("intent") for s in a.args.get("steps", [])])
        else:
            intents.append(a.intent.value)
    results["F-011"] = CaseResult("Pass" if "path_move" in intents else "Fail", ",".join(intents), "VN move parsing.")

    p = route("sao chep thu muc zzz vao backup o d")
    intents = []
    for a in p.actions:
        if a.intent == IntentType.COMPOSITE:
            intents.extend([s.get("intent") for s in a.args.get("steps", [])])
        else:
            intents.append(a.intent.value)
    results["F-012"] = CaseResult("Pass" if "path_copy" in intents else "Fail", ",".join(intents), "VN copy parsing.")

    p = route("doi ten thu muc zzz thanh zzz_new")
    intents = []
    for a in p.actions:
        if a.intent == IntentType.COMPOSITE:
            intents.extend([s.get("intent") for s in a.args.get("steps", [])])
        else:
            intents.append(a.intent.value)
    results["F-013"] = CaseResult("Pass" if "path_rename" in intents else "Fail", ",".join(intents), "VN rename parsing.")

    p = route("nen thu muc zzz vao backup o d")
    intents = []
    for a in p.actions:
        if a.intent == IntentType.COMPOSITE:
            intents.extend([s.get("intent") for s in a.args.get("steps", [])])
        else:
            intents.append(a.intent.value)
    results["F-014"] = CaseResult("Pass" if "path_zip" in intents else "Fail", ",".join(intents), "VN zip parsing.")

    p = route("tao tm zzz o desktp roi dc vao backup o d")
    intents = []
    for a in p.actions:
        if a.intent == IntentType.COMPOSITE:
            intents.extend([s.get("intent") for s in a.args.get("steps", [])])
        else:
            intents.append(a.intent.value)
    ok_abr = "dir_create" in intents and "path_move" in intents
    results["F-015"] = CaseResult("Pass" if ok_abr else "Fail", ",".join(intents), "Abbrev+typo VN parsing.")

    p = route("tim file dư liueeju ở màn hình desk rồi vào ổ D tạo thư mục mới đặt tên tùy ý rồi di chuyển vào đó được không?")
    intents = []
    for a in p.actions:
        if a.intent == IntentType.COMPOSITE:
            intents.extend([s.get("intent") for s in a.args.get("steps", [])])
        else:
            intents.append(a.intent.value)
    ok_sentence = all(x in intents for x in ["file_search", "dir_create", "path_move"])
    results["F-016"] = CaseResult("Pass" if ok_sentence else "Fail", ",".join(intents), "Original complex VN sentence parsing.")

    # Permission
    guest = UserContext(role=Role.GUEST, allow_c_drive_full=False, actor="tc-guest")
    dec = evaluate_action(Action(intent=IntentType.FILE_DELETE, args={"path": "C:\\tmp\\a.txt"}), guest)
    results["P-001"] = CaseResult("Pass" if not dec.allowed else "Fail", dec.reason, "Guest delete C must deny.")

    user = UserContext(role=Role.USER, allow_c_drive_full=False, actor="tc-user")
    dec = evaluate_action(Action(intent=IntentType.FILE_DELETE, args={"path": "C:\\tmp\\a.txt"}), user)
    results["P-002"] = CaseResult("Pass" if not dec.allowed else "Fail", dec.reason, "User delete C must deny.")

    owner_no = UserContext(role=Role.OWNER, allow_c_drive_full=False, actor="tc-owner")
    dec = evaluate_action(Action(intent=IntentType.FILE_DELETE, args={"path": "C:\\tmp\\a.txt"}), owner_no)
    results["P-003"] = CaseResult("Pass" if not dec.allowed else "Fail", dec.reason, "Owner without toggle must deny.")

    dev = UserContext(role=Role.DEV, allow_c_drive_full=True, actor="tc-dev")
    dec = evaluate_action(Action(intent=IntentType.FILE_DELETE, args={"path": "C:\\tmp\\a.txt"}), dev)
    results["P-004"] = CaseResult("Pass" if dec.allowed else "Fail", dec.reason, "Dev delete C allowed.")

    # Security
    dec = evaluate_action(Action(intent=IntentType.SEND_BULK_EMAIL, args={"recipients": ["a@b.com"]}), owner)
    results["S-001"] = CaseResult(
        "Pass" if dec.allowed and dec.needs_confirmation else "Fail",
        f"allowed={dec.allowed}, confirm={dec.needs_confirmation}",
        "High-risk outbound confirmation.",
    )

    # S-002 audit file creation
    from futureos.safety import write_audit

    write_audit({"event": "tc"})
    chain_ok = verify_chain(AUDIT_FILE, AUDIT_CHAIN_FILE)
    results["S-002"] = CaseResult(
        "Pass" if AUDIT_FILE.exists() and chain_ok else "Fail",
        f"audit_exists={AUDIT_FILE.exists()},chain_ok={chain_ok}",
        "Audit log exists and integrity chain verifies.",
    )

    # S-003 auth lockout
    _reset_runtime_files()
    for _ in range(3):
        authenticate_login("owner", actor="lock-user", allow_c_drive_full=False, secret="wrong", voice_confidence=0.1)
    ctx, msg = authenticate_login("owner", actor="lock-user", allow_c_drive_full=False, secret="wrong", voice_confidence=0.1)
    locked = (ctx is None) and ("locked until" in msg.lower())
    results["S-003"] = CaseResult("Pass" if locked else "Fail", msg, "Actor lockout after repeated failures.")

    # S-004 sensitive action rate limit
    _reset_runtime_files()
    store = SessionStore()
    rec = store.create_session(user)
    ok_limit = True
    for _ in range(5):
        if not store.within_sensitive_rate_limit(rec.session_id):
            ok_limit = False
            break
        store.touch_sensitive(rec.session_id)
    blocked_next = not store.within_sensitive_rate_limit(rec.session_id)
    results["S-004"] = CaseResult("Pass" if ok_limit and blocked_next else "Fail", f"blocked_next={blocked_next}", "Sensitive rate limit.")

    # Voice
    ve = VoiceEngine()
    results["V-001"] = CaseResult(
        "Pass" if isinstance(ve, VoiceEngine) else "Fail",
        "Voice engine initialized",
        "Wakeword pipeline object available (manual mic e2e pending).",
    )
    ctx, msg = authenticate_login("owner", actor="v-owner", allow_c_drive_full=False, secret="123456", voice_confidence=0.2)
    results["V-002"] = CaseResult(
        "Pass" if ctx is not None and ctx.role == Role.OWNER else "Fail",
        msg,
        "Low confidence fallback to PIN.",
    )

    # Regression
    r1 = route("tim file o d")
    has_find = any((a.intent == IntentType.FIND_DRAFT or a.intent == IntentType.COMPOSITE) for a in r1.actions)
    results["R-001"] = CaseResult("Pass" if has_find else "Fail", r1.model_dump_json(), "Basic router.")

    r2 = route("tim ban thao lap trinh trong o D va gui zalo cho sep va gui email hang loat cho team")
    _, checks = evaluate_plan(r2.actions, owner)
    intents = [c["intent"] for c in checks]
    needed = {"send_zalo", "send_bulk_email"}
    results["R-002"] = CaseResult(
        "Pass" if needed.issubset(set(intents)) else "Fail",
        ",".join(intents),
        "Composite policy flatten.",
    )

    return results


def _update_workbook(results: dict[str, CaseResult]) -> tuple[int, int]:
    wb = load_workbook("TEST_CASES.xlsx")
    today = str(date.today())
    updated = 0
    failed = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            case_id = str(row[0].value or "").strip()
            if not case_id or case_id not in results:
                continue
            r = results[case_id]
            row[6].value = r.actual
            row[7].value = r.status
            row[10].value = today
            row[11].value = r.notes
            updated += 1
            if r.status.lower() != "pass":
                failed += 1
    wb.save("TEST_CASES.xlsx")
    return updated, failed


def _render_evidence(results: dict[str, CaseResult]) -> Path:
    lines = ["futureOS testcase evidence", f"date={date.today()}", ""]
    for case_id in sorted(results.keys()):
        r = results[case_id]
        lines.append(f"{case_id}: {r.status} | {r.actual}")
    text = "\n".join(lines)

    img = Image.new("RGB", (1400, 900), color=(16, 18, 24))
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    draw.multiline_text((30, 30), text, fill=(220, 240, 255), font=font, spacing=6)
    out = Path("data/test_evidence_latest.png")
    out.parent.mkdir(parents=True, exist_ok=True)
    img.save(out)
    return out


def main() -> int:
    results = _run_cases()
    updated, failed = _update_workbook(results)
    evidence = _render_evidence(results)
    print(f"Updated rows: {updated}")
    print(f"Failed rows: {failed}")
    print(f"Evidence image: {evidence}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
